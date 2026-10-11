"""Make text copies of a course folder into the parallel _text/ tree.

Usage: _system/venv/bin/python _system/tools/make_text_copies.py "<course folder>"

- PDF  -> <name>.pdf.txt   (PDFKit text layer + Vision OCR, via _system/tools/pdf2text.js)
- DOCX -> <name>.docx.txt  (MarkItDown: headings, lists, tables as Markdown; textutil as fallback)
- PPTX -> <name>.pptx.txt  (MarkItDown with "=== Slide N ===" markers; python-pptx as fallback)
- XLSX -> <name>.xlsx - <sheet>.csv  (openpyxl; one CSV per sheet, computed values)

Originals are only read, never changed. An existing text copy is left alone unless
it is stale (the original was modified after the copy was made): a stale copy is
regenerated and reported as REGENERATED. A report line is printed per file.
"""
import csv
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches  # noqa: F401  (confirms python-pptx import)

ROOT = Path(__file__).resolve().parents[2]
PDF2TEXT = ["/usr/bin/osascript", "-l", "JavaScript", str(ROOT / "_system" / "tools" / "pdf2text.js")]


def clean(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "-", name).strip() or "sheet"


def stale(src: Path, dst: Path) -> bool:
    """A text copy is stale when its original was modified after the copy was made."""
    return dst.exists() and src.stat().st_mtime > dst.stat().st_mtime


def markitdown_text(src: Path) -> str:
    """Convert one local file with MarkItDown (local, offline; no plugins, no LLM, no cloud)."""
    from markitdown import MarkItDown
    return MarkItDown(enable_plugins=False).convert_local(str(src)).markdown


def markitdown_pptx(src: Path, dst: Path) -> str:
    """MarkItDown text of a deck, with its slide comments turned into our "=== Slide N ===" markers."""
    text = markitdown_text(src)
    text = re.sub(r"^<!-- Slide number: (\d+) -->$", r"=== Slide \1 ===", text, flags=re.M)
    text = re.sub(r"^!\[[^\]]*\]\([^)]*\)\s*$\n?", "", text, flags=re.M)  # image links carry no text
    n = len(re.findall(r"^=== Slide \d+ ===$", text, flags=re.M))
    if not n:
        raise ValueError("no slide markers in MarkItDown output")
    dst.write_text(text, encoding="utf-8")
    return f"slides={n} (markitdown)"


def pptx_text(src: Path, dst: Path) -> str:
    prs = Presentation(str(src))
    lines, no_text = [], []
    for n, slide in enumerate(prs.slides, 1):
        lines.append(f"=== Slide {n} ===")
        got = False

        def walk(shapes):
            nonlocal got
            for sh in shapes:
                if sh.shape_type == 6:  # group
                    walk(sh.shapes)
                    continue
                if sh.has_text_frame and sh.text_frame.text.strip():
                    lines.append(sh.text_frame.text.strip())
                    got = True
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        lines.append(" | ".join(c.text.strip() for c in row.cells))
                    got = True

        walk(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("[Notes]")
                lines.append(notes)
        if not got:
            no_text.append(n)
        lines.append("")
    dst.write_text("\n".join(lines), encoding="utf-8")
    missing = ",".join(map(str, no_text)) if no_text else "-"
    return f"slides={len(prs.slides)} no_text={missing}"


def xlsx_csv(src: Path, dst_dir: Path) -> str:
    wb = load_workbook(str(src), data_only=True)
    made = []
    for ws in wb.worksheets:
        out = dst_dir / f"{src.name} - {clean(ws.title)}.csv"
        redo = stale(src, out)
        if out.exists() and not redo:
            made.append(f"{ws.title}(exists)")
            continue
        with out.open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            for row in ws.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])
        made.append(f"{ws.title}({ws.max_row}x{ws.max_column}{', REGENERATED' if redo else ''})")
    current = {f"{src.name} - {clean(ws.title)}.csv" for ws in wb.worksheets}
    leftover = sorted(p.name for p in dst_dir.glob("*.csv")
                      if p.name.startswith(src.name + " - ") and p.name not in current)
    note = f" | leftover CSVs for sheets no longer in the workbook (kept): {', '.join(leftover)}" if leftover else ""
    return f"sheets={len(wb.worksheets)} " + "; ".join(made) + note


def main() -> None:
    course = Path(sys.argv[1])
    src_root = ROOT / course
    dst_root = ROOT / "_text" / course
    for src in sorted(p for p in src_root.rglob("*") if p.is_file()):
        ext = src.suffix.lower()
        if ext not in (".pdf", ".docx", ".pptx", ".xlsx"):
            continue
        if src.name.startswith("~$"):  # Office lock file, not a document
            continue
        if re.search(r"mylab|pearson|mastering", src.name, re.I):  # publisher content: never read (CLAUDE.md)
            print(f"{src.relative_to(src_root)} | SKIPPED (Pearson/MyLab item, not read)", flush=True)
            continue
        rel = src.relative_to(src_root)
        dst_dir = dst_root / rel.parent
        dst_dir.mkdir(parents=True, exist_ok=True)
        try:
            if ext == ".xlsx":
                info = xlsx_csv(src, dst_dir)
            else:
                dst = dst_dir / (src.name + ".txt")
                redo = stale(src, dst)
                if dst.exists() and not redo:
                    info = "exists, skipped"
                elif ext == ".pdf":
                    r = subprocess.run(PDF2TEXT + [str(src), str(dst)], capture_output=True, text=True)
                    info = (r.stdout.strip().splitlines() or [r.stderr.strip() or "no output"])[-1]
                elif ext == ".docx":
                    try:
                        dst.write_text(markitdown_text(src), encoding="utf-8")
                        info = "ok (markitdown)"
                    except Exception as exc:
                        info = f"markitdown failed ({type(exc).__name__}), textutil used: "
                        r = subprocess.run(["/usr/bin/textutil", "-convert", "txt", "-encoding", "UTF-8",
                                            "-output", str(dst), str(src)], capture_output=True, text=True)
                        info += "ok" if r.returncode == 0 else f"ERROR {r.stderr.strip()}"
                else:
                    try:
                        info = markitdown_pptx(src, dst)
                    except Exception as exc:
                        info = f"markitdown failed ({type(exc).__name__}), python-pptx used: " + pptx_text(src, dst)
                if redo and not info.startswith("ERROR"):
                    info = "REGENERATED (original changed) | " + info
        except Exception as exc:  # report and continue with the next file
            info = f"ERROR {type(exc).__name__}: {exc}"
        print(f"{rel} | {info}", flush=True)


if __name__ == "__main__":
    main()
